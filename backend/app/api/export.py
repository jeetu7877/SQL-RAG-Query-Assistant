"""Database export endpoints for CSV and Excel downloads."""
import csv
import io
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from starlette.background import BackgroundTask
from fastapi.responses import FileResponse
from sqlalchemy import inspect, text
from sqlalchemy.engine import Engine

from app.api.deps import get_connection
from app.database.connection_manager import ActiveConnection

router = APIRouter(prefix="/api/database", tags=["export"])

MAX_EXPORT_ROWS = 100_000


def _safe_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._")
    return value or "database"


def _cleanup(path: str) -> None:
    Path(path).unlink(missing_ok=True)


def _jsonable(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    if isinstance(value, (bytes, bytearray, memoryview)):
        return "<binary>"
    return str(value)


def _table_names(engine: Engine) -> list[str]:
    return sorted(inspect(engine).get_table_names(schema="public"))


def _read_table(engine: Engine, table: str) -> tuple[list[str], list[list[Any]]]:
    # The table name comes from SQLAlchemy's inspected public-schema table list,
    # but quote it anyway because identifiers cannot be parameterized.
    sql = text(f'SELECT * FROM "{table.replace(chr(34), chr(34) * 2)}" LIMIT {MAX_EXPORT_ROWS + 1}')
    with engine.connect() as conn:
        result = conn.execute(sql)
        columns = list(result.keys())
        rows = [[_jsonable(v) for v in row] for row in result.fetchmany(MAX_EXPORT_ROWS + 1)]
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(
            status_code=413,
            detail=f'Table "{table}" has more than {MAX_EXPORT_ROWS:,} rows. Export a filtered result from Chat instead.',
        )
    return columns, rows


def _csv_bytes(columns: list[str], rows: list[list[Any]]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(columns)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8-sig")


def _excel_file(engine: Engine, tables: list[str], path: str) -> None:
    from openpyxl import Workbook

    print("EXCEL EXPORT START")
    print("TABLES TO EXPORT:", tables)

    # Normal workbook instead of write_only workbook
    wb = Workbook()

    # Remove the default "Sheet"
    default_sheet = wb.active
    wb.remove(default_sheet)

    used_sheet_names: set[str] = set()

    for table in tables:
        print(f"EXPORTING TABLE: {table}")

        columns, rows = _read_table(engine, table)

        # Excel sheet names:
        # - maximum 31 characters
        # - cannot contain \ / * ? : [ ]
        base = re.sub(r'[\\/*?:\[\]]', "_", table)
        base = base[:31] or "Sheet"

        sheet_name = base
        suffix = 1

        while sheet_name in used_sheet_names:
            tail = f"_{suffix}"
            sheet_name = f"{base[:31 - len(tail)]}{tail}"
            suffix += 1

        used_sheet_names.add(sheet_name)

        print(f"CREATING SHEET: {sheet_name}")

        ws = wb.create_sheet(title=sheet_name)

        # Header
        ws.append(columns)

        # Data
        for row in rows:
            ws.append(row)

    # If database has no tables
    if not tables:
        ws = wb.create_sheet(title="Database")
        ws.append(["No tables found in public schema"])

    print("FINAL EXCEL SHEETS:", wb.sheetnames)

    wb.save(path)

    print("EXCEL FILE SAVED:", path)

@router.get("/export")
def export_database(
    format: str = Query("xlsx", pattern="^(csv|xlsx)$"),
    table: str | None = Query(default=None),
    active: ActiveConnection = Depends(get_connection),
):
    tables = _table_names(active.engine)
    print("EXPORT TABLES:", tables)
    print("REQUESTED TABLE:", table)
    if table is not None and table not in tables:
        raise HTTPException(status_code=404, detail=f'Table "{table}" was not found in the public schema.')

    selected = [table] if table else tables
    db_name = _safe_filename(active.database_name)

    if format == "csv":
        if table:
            columns, rows = _read_table(active.engine, table)
            data = _csv_bytes(columns, rows)
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".csv")
            tmp_path = tmp.name
            try:
                tmp.write(data)
                tmp.close()
            except Exception:
                tmp.close()
                Path(tmp_path).unlink(missing_ok=True)
                raise
            return FileResponse(
                tmp_path,
                media_type="text/csv",
                filename=f"{_safe_filename(table)}.csv",
                background=BackgroundTask(_cleanup, tmp_path),
            )

        # A database can contain multiple tables, so a full CSV export is a ZIP
        # containing one CSV file per table.
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".zip")
        tmp_path = tmp.name
        tmp.close()
        try:
            with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as archive:
                for name in selected:
                    columns, rows = _read_table(active.engine, name)
                    archive.writestr(f"{_safe_filename(name)}.csv", _csv_bytes(columns, rows))
        except Exception:
            Path(tmp_path).unlink(missing_ok=True)
            raise
        return FileResponse(
            tmp_path,
            media_type="application/zip",
            filename=f"{db_name}_csv_export.zip",
            background=BackgroundTask(_cleanup, tmp_path),
        )

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp_path = tmp.name
    tmp.close()
    try:
        _excel_file(active.engine, selected, tmp_path)
    except Exception:
        Path(tmp_path).unlink(missing_ok=True)
        raise

    return FileResponse(
        tmp_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{db_name}_export.xlsx" if not table else f"{_safe_filename(table)}.xlsx",
        background=BackgroundTask(_cleanup, tmp_path),
    )
